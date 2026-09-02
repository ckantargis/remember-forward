"""
Render a .docx to a matching .pdf.

WHY THIS EXISTS
  Every document in this project is published in both formats. The 2026-08-20
  errata edited .docx files and left the .pdf copies untouched, so archive.org
  carried corrected Word files beside stale PDFs of the same document. There is
  no Word or LibreOffice on this machine to convert with, so this reads the
  docx structure directly and rebuilds it with reportlab.

  It is not a pixel-faithful converter. It preserves heading level, body text,
  bullets, tables and the centred title block — which is all these flyers use.

    from docx_to_pdf import convert
    convert("containers/gd_10_flyer_titanium.docx")
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (ListFlowable, ListItem, Paragraph, SimpleDocTemplate,
                                Spacer, Table, TableStyle)

FONT = "Helvetica"
ACCENT = colors.HexColor("#1F3B53")


def _s(name, size, leading, bold=False, italic=False, center=False, **kw):
    f = FONT + ("-BoldOblique" if bold and italic else "-Bold" if bold
                else "-Oblique" if italic else "")
    return ParagraphStyle(name, fontName=f, fontSize=size, leading=leading,
                          alignment=TA_CENTER if center else 0, **kw)


STYLES = {
    "title":  _s("title", 26, 30, bold=True, center=True, spaceAfter=2),
    "sub":    _s("sub", 17, 21, bold=True, center=True, spaceAfter=4),
    "sub2":   _s("sub2", 11, 14, italic=True, center=True, spaceAfter=6),
    "strip":  _s("strip", 11, 14, bold=True, center=True, spaceAfter=12, textColor=ACCENT),
    "h1":     _s("h1", 13, 16, bold=True, spaceBefore=14, spaceAfter=5, textColor=ACCENT),
    "h2":     _s("h2", 11, 14, bold=True, spaceBefore=9, spaceAfter=3),
    "body":   _s("body", 10.5, 14, spaceAfter=6),
    "italic": _s("italic", 10, 13.5, italic=True, spaceAfter=4),
    "li":     _s("li", 10.5, 13.5, spaceAfter=2),
    "cell":   _s("cell", 8.5, 11),
    "cellh":  _s("cellh", 8.5, 11, bold=True),
}


def _esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _style_for(p, index):
    """Map a docx paragraph to one of our PDF styles."""
    name = (p.style.name if p.style is not None else "") or ""
    runs = p.runs
    size = next((r.font.size.pt for r in runs if r.font.size), None)
    bold = bool(runs and runs[0].font.bold)
    italic = bool(runs and runs[0].font.italic)
    centered = p.alignment == WD_ALIGN_PARAGRAPH.CENTER

    if centered and size and size >= 20:
        return "title"
    if centered and size and size >= 15:
        return "sub"
    if centered and italic:
        return "sub2"
    if centered and bold:
        return "strip"
    if "Heading 1" in name or (bold and size and 12 <= size < 15):
        return "h1"
    if "Heading 2" in name or (bold and size and 10.5 <= size < 12):
        return "h2"
    if "List" in name:
        return "li"
    if italic:
        return "italic"
    return "body"


def convert(docx_path, pdf_path=None):
    pdf_path = pdf_path or os.path.splitext(docx_path)[0] + ".pdf"
    d = Document(docx_path)
    story, pending = [], []

    def flush():
        if pending:
            story.append(ListFlowable(
                [ListItem(Paragraph(_esc(t), STYLES["li"]), leftIndent=14) for t in pending],
                bulletType="bullet", start="•", leftIndent=14, bulletFontSize=8))
            story.append(Spacer(1, 4))
            pending.clear()

    # docx keeps paragraphs and tables in separate lists; interleave by body order
    body = d.element.body
    p_iter, t_iter = iter(d.paragraphs), iter(d.tables)
    idx = 0
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = next(p_iter, None)
            if p is None or not p.text.strip():
                continue
            st = _style_for(p, idx)
            idx += 1
            if st == "li":
                pending.append(p.text)
            else:
                flush()
                story.append(Paragraph(_esc(p.text), STYLES[st]))
        elif tag == "tbl":
            flush()
            t = next(t_iter, None)
            if t is None:
                continue
            rows = [[c.text for c in r.cells] for r in t.rows]
            if not rows:
                continue
            ncol = len(rows[0])
            avail = LETTER[0] - 2 * 54
            widths = ([avail * w for w in (.22, .13, .30, .35)] if ncol == 4
                      else [avail * w for w in (.46, .16, .38)] if ncol == 3
                      else [avail / ncol] * ncol)
            data = [[Paragraph(_esc(c), STYLES["cellh"] if i == 0 else STYLES["cell"])
                     for c in row] for i, row in enumerate(rows)]
            tb = Table(data, colWidths=widths, repeatRows=1)
            tb.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF2")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story += [tb, Spacer(1, 6)]
    flush()

    SimpleDocTemplate(pdf_path, pagesize=LETTER, leftMargin=54, rightMargin=54,
                      topMargin=54, bottomMargin=54,
                      title=os.path.basename(pdf_path)).build(story)
    return pdf_path


if __name__ == "__main__":
    import sys
    for f in sys.argv[1:]:
        print("  ->", convert(f))
